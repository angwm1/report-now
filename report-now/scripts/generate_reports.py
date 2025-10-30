from pathlib import Path

from docx import Document


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw in lines:
        stripped = raw.strip()
        if not stripped:
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        rows.append(cells)

    cleaned: list[list[str]] = []
    for row in rows:
        if all(set(cell) <= set("-: ") for cell in row):
            continue
        cleaned.append(row)

    if not cleaned:
        return []

    col_count = max(len(row) for row in cleaned)
    for row in cleaned:
        if len(row) < col_count:
            row.extend([""] * (col_count - len(row)))
    return cleaned


def md_to_docx(md_path: Path, docx_path: Path, title: str) -> None:
    doc = Document()
    doc.add_heading(title, level=0)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        if not stripped:
            doc.add_paragraph("")
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = stripped[level:].strip()
            level = max(1, min(level, 4))
            doc.add_heading(heading, level=level)
            i += 1
            continue

        if stripped.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                item = lines[i].strip()[2:].strip()
                doc.add_paragraph(item, style="List Bullet")
                i += 1
            continue

        if "|" in stripped:
            table_block: list[str] = []
            while i < len(lines) and "|" in lines[i]:
                table_block.append(lines[i])
                i += 1

            table_rows = parse_table(table_block)
            if table_rows:
                table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                table.style = "Light Grid Accent 2"
                for r_idx, row in enumerate(table_rows):
                    for c_idx, cell in enumerate(row):
                        table.cell(r_idx, c_idx).text = cell
            continue

        doc.add_paragraph(stripped)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def generate_reports() -> None:
    base = Path("tests")
    output = Path("reports")
    docs = [
        ("test-plan.md", "ReportNow IEEE 829 Test Plan"),
        ("test-cases.md", "ReportNow Test Cases Catalogue"),
        ("requirements-coverage.md", "ReportNow Requirements Coverage Matrix"),
    ]

    for md_file, title in docs:
        md_path = base / md_file
        docx_path = output / (title.replace(" ", "_") + ".docx")
        md_to_docx(md_path, docx_path, title)
        print(f"Generated {docx_path}")


if __name__ == "__main__":
    generate_reports()
